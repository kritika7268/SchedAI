import csv
import io
from datetime import date, datetime, time as dtime, timedelta
from typing import Literal, Optional
from xml.sax.saxutils import escape

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.concurrency import run_in_threadpool
from fastapi.responses import StreamingResponse
from ortools.sat.python import cp_model
from pydantic import BaseModel

from app.database import get_db_connection
from app.auth import get_current_user, require_role
from app.websocket_manager import manager

router = APIRouter(prefix="/exams", tags=["Exams"])

# ------------------------------------------------------------
# SETTINGS (edit here)
# ------------------------------------------------------------
INTERNAL_ROUNDS = {
    "class_test_1": "Class Test 1",
    "mid_sem": "Mid Sem",
    "class_test_2": "Class Test 2",
}
ROUND_LABELS = {**INTERNAL_ROUNDS, "external": "External Exam"}

# Printed at the top of the Excel / PDF downloads
DEPARTMENT_NAME = "Department of Computer Applications"

# Internal papers are 1 hour each. A batch with two papers in a day
# writes one at 10:00 and one at 11:30.
SESSIONS = [
    {"no": 1, "start": "10:00:00", "end": "11:00:00"},
    {"no": 2, "start": "11:30:00", "end": "12:30:00"},
]

# Used when the external file has no start_time / end_time
DEFAULT_EXT_START = dtime(10, 0)
DEFAULT_EXT_END = dtime(13, 0)

SOLVER_SECONDS = 20
EXTERNAL_COLUMNS = ["batch_name", "subject_code", "exam_date", "start_time", "end_time", "room_name"]
EXTERNAL_REQUIRED = ["batch_name", "subject_code", "exam_date"]


class GenerateIn(BaseModel):
    exam_round: Literal["class_test_1", "mid_sem", "class_test_2"]
    week_start: date                      # must be a Monday
    ignore_holidays: bool = False


class InvigilatorIn(BaseModel):
    teacher_id: int


# ------------------------------------------------------------
# helpers
# ------------------------------------------------------------
async def _notify(title: str, message: str):
    conn = get_db_connection()
    cur = conn.cursor()
    try:
        cur.execute(
            "INSERT INTO notifications (title, message, notif_type, audience) "
            "VALUES (%s, %s, %s, %s)", (title, message, "info", "all"))
        conn.commit()
        nid = cur.lastrowid
    finally:
        cur.close()
        conn.close()
    await manager.broadcast({
        "notification_id": nid, "title": title, "message": message,
        "notif_type": "info", "audience": "all",
        "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    })


def _hhmm(v):
    if isinstance(v, timedelta):
        s = int(v.total_seconds())
        return f"{s // 3600:02d}:{(s % 3600) // 60:02d}"
    return str(v)[:5]


def _as_date(v):
    return v.date() if isinstance(v, datetime) else v


def _mins(hms: str) -> int:
    p = hms.split(":")
    return int(p[0]) * 60 + int(p[1])


def _tmins(v):
    if isinstance(v, timedelta):
        return int(v.total_seconds() // 60)
    p = str(v).split(":")
    return int(p[0]) * 60 + int(p[1])


def _overlap(s1, e1, s2, e2) -> bool:
    return s1 < e2 and s2 < e1


def exam_week_containing(cur, d: date):
    """{'week_start','week_end'} if d falls in a week that has exams, else None.
    `cur` must be a dictionary cursor."""
    cur.execute(
        "SELECT DISTINCT DATE_SUB(exam_date, INTERVAL WEEKDAY(exam_date) DAY) AS wk "
        "FROM exam_schedule")
    for r in cur.fetchall():
        wk = _as_date(r["wk"])
        if wk <= d <= wk + timedelta(days=4):
            return {"week_start": wk, "week_end": wk + timedelta(days=4)}
    return None


def _holiday_days(cur, week_start: date, week_end: date):
    cur.execute(
        "SELECT holiday_date, end_date FROM holidays "
        "WHERE holiday_date <= %s AND COALESCE(end_date, holiday_date) >= %s",
        (week_end, week_start))
    days = set()
    for h in cur.fetchall():
        d = max(_as_date(h["holiday_date"]), week_start)
        end = min(_as_date(h["end_date"] or h["holiday_date"]), week_end)
        while d <= end:
            days.add(d)
            d += timedelta(days=1)
    return days


def _assign_invigilators(items, teachers, absent, own):
    """
    items:   [{"date": date, "start": minutes, "end": minutes, "subject_id": int}]
    teachers:{teacher_id: name}   absent: {(teacher_id, date)}   own: {subject_id: {teacher_id}}
    One invigilator per item. Nobody works two overlapping papers.
    Duties are shared as equally as possible (min max-min difference).
    Returns (list of teacher_id per item, list of teachers who could take duty).
    """
    tids = list(teachers)
    N = len(items)
    m = cp_model.CpModel()
    w = {}
    for i, it in enumerate(items):
        ok = [t for t in tids if (t, it["date"]) not in absent]
        if not ok:
            raise HTTPException(400, f"No teacher is available on {it['date']}.")
        for t in ok:
            w[i, t] = m.NewBoolVar(f"w_{i}_{t}")
        m.Add(sum(w[i, t] for t in ok) == 1)

    by_date = {}
    for i, it in enumerate(items):
        by_date.setdefault(it["date"], []).append(i)

    for idxs in by_date.values():
        for a in range(len(idxs)):
            for b in range(a + 1, len(idxs)):
                i, j = idxs[a], idxs[b]
                if _overlap(items[i]["start"], items[i]["end"], items[j]["start"], items[j]["end"]):
                    for t in tids:
                        if (i, t) in w and (j, t) in w:
                            m.Add(w[i, t] + w[j, t] <= 1)

    eligible = [t for t in tids if any((i, t) in w for i in range(N))]
    mx = m.NewIntVar(0, N, "mx")
    mn = m.NewIntVar(0, N, "mn")
    for t in eligible:
        duty_t = sum(w[i, t] for i in range(N) if (i, t) in w)
        m.Add(duty_t <= mx)
        m.Add(duty_t >= mn)

    extras = []
    for t in tids:
        for d, idxs in by_date.items():
            vs = [w[i, t] for i in idxs if (i, t) in w]
            if len(vs) > 1:
                e = m.NewBoolVar(f"e_{t}_{d}")
                m.Add(sum(vs) - 1 <= e)        # two papers on one day = extra
                extras.append(e)
    own_pen = sum(w[i, t] for (i, t) in w if t in own.get(items[i]["subject_id"], ()))
    m.Minimize(1000 * (mx - mn) + 10 * sum(extras) + own_pen)

    s = cp_model.CpSolver()
    s.parameters.max_time_in_seconds = SOLVER_SECONDS
    s.parameters.num_workers = 4
    st = s.Solve(m)
    if st not in (cp_model.OPTIMAL, cp_model.FEASIBLE):
        raise HTTPException(
            400, "Could not assign invigilators. There are not enough free teachers for papers "
                 "held at the same time (check teacher absences).")

    result = [None] * N
    for (i, t), var in w.items():
        if s.Value(var):
            result[i] = t
    return result, eligible


def _fairness(duty, teachers, eligible):
    el = [duty[t] for t in eligible]
    return {
        "min_duties": min(el) if el else 0,
        "max_duties": max(el) if el else 0,
        "teachers": sorted(
            ({"teacher_name": teachers[t], "duties": duty[t]} for t in teachers),
            key=lambda r: (-r["duties"], r["teacher_name"])),
    }


def _load_teachers_and_own(cur):
    cur.execute("SELECT * FROM teachers ORDER BY teacher_id")
    teachers = {t["teacher_id"]: t["teacher_name"]
                for t in cur.fetchall() if t.get("status") != "inactive"}
    if not teachers:
        raise HTTPException(400, "No active teachers found.")
    own = {}
    cur.execute("SELECT teacher_id, subject_id FROM teacher_skills")
    for r in cur.fetchall():
        own.setdefault(r["subject_id"], set()).add(r["teacher_id"])
    return teachers, own


def _load_rooms(cur):
    cur.execute("SELECT * FROM rooms")
    rooms = []
    for r in cur.fetchall():
        if r.get("is_available", 1) not in (1, True, None):
            continue
        rooms.append({
            "room_id": r["room_id"], "room_name": r["room_name"],
            "room_type": r.get("room_type") or "",
            "capacity": int(r.get("capacity") or 0),
            "is_lab": "lab" in str(r.get("room_type") or "").lower(),
        })
    if not rooms:
        raise HTTPException(400, "No rooms available for exams.")
    return rooms


# ------------------------------------------------------------
# CONFLICT CHECK
# ------------------------------------------------------------
def _find_conflicts(cur, exam_round=None):
    """Checks the saved schedule. `cur` must be a dictionary cursor."""
    cur.execute(
        """SELECT e.exam_id, e.exam_round, e.exam_date, e.start_time, e.end_time,
                  e.batch_id, b.batch_name, e.room_id, r.room_name, r.capacity,
                  e.invigilator_id, t.teacher_name, s.subject_name
           FROM exam_schedule e
           JOIN batches b ON e.batch_id = b.batch_id
           JOIN rooms r ON e.room_id = r.room_id
           JOIN teachers t ON e.invigilator_id = t.teacher_id
           JOIN subjects s ON e.subject_id = s.subject_id""")
    rows = cur.fetchall()
    for r in rows:
        r["date"] = _as_date(r["exam_date"])
        r["s"] = _tmins(r["start_time"])
        r["e"] = _tmins(r["end_time"])

    cur.execute("SELECT batch_id, COUNT(*) AS c FROM students GROUP BY batch_id")
    counts = {x["batch_id"]: x["c"] for x in cur.fetchall()}
    cur.execute("SELECT teacher_id, absence_date FROM teacher_absences")
    absent = {(a["teacher_id"], _as_date(a["absence_date"])) for a in cur.fetchall()}

    scope = [r for r in rows if exam_round is None or r["exam_round"] == exam_round]
    scope_ids = {r["exam_id"] for r in scope}
    found = []

    def tag(r):
        return f"{r['subject_name']} - {r['batch_name']} ({r['date']}, {_hhmm(r['start_time'])})"

    for i in range(len(rows)):
        for j in range(i + 1, len(rows)):
            a, b = rows[i], rows[j]
            if a["date"] != b["date"] or not _overlap(a["s"], a["e"], b["s"], b["e"]):
                continue
            if a["exam_id"] not in scope_ids and b["exam_id"] not in scope_ids:
                continue
            if a["batch_id"] == b["batch_id"]:
                found.append({"type": "Batch clash",
                              "detail": f"{a['batch_name']} has two exams at the same time: {tag(a)} and {tag(b)}"})
            if a["room_id"] == b["room_id"]:
                found.append({"type": "Room clash",
                              "detail": f"{a['room_name']} is used twice at the same time: {tag(a)} and {tag(b)}"})
            if a["invigilator_id"] == b["invigilator_id"]:
                found.append({"type": "Invigilator clash",
                              "detail": f"{a['teacher_name']} has two duties at the same time: {tag(a)} and {tag(b)}"})

    for r in scope:
        if (r["invigilator_id"], r["date"]) in absent:
            found.append({"type": "Invigilator absent",
                          "detail": f"{r['teacher_name']} is marked absent on {r['date']} but has duty: {tag(r)}"})
        size = counts.get(r["batch_id"], 0)
        if size > int(r["capacity"] or 0):
            found.append({"type": "Room too small",
                          "detail": f"{r['room_name']} holds {r['capacity']} but {r['batch_name']} has {size} students: {tag(r)}"})
        if r["date"].weekday() == 6:
            found.append({"type": "Sunday exam", "detail": f"Exam is on a Sunday: {tag(r)}"})

    return {"checked": len(scope), "conflicts": found}


# ------------------------------------------------------------
# INTERNAL: generate (all batches, one click, one week)
# ------------------------------------------------------------
def _generate_sync(body: GenerateIn):
    ws = body.week_start
    if ws.weekday() != 0:
        raise HTTPException(400, "Please choose a Monday as the first day of the exam week.")
    we = ws + timedelta(days=4)
    label = INTERNAL_ROUNDS[body.exam_round]

    conn = get_db_connection()
    cur = conn.cursor(dictionary=True)
    try:
        # ---- exam days (Mon-Fri minus holidays) ----
        holidays = set() if body.ignore_holidays else _holiday_days(cur, ws, we)
        days = [ws + timedelta(days=i) for i in range(5)
                if (ws + timedelta(days=i)) not in holidays]
        skipped_days = [d.isoformat() for d in sorted(holidays)]
        if not days:
            raise HTTPException(400, "Every day of that week is a holiday.")
        D = len(days)

        teachers, own = _load_teachers_and_own(cur)
        cur.execute(
            "SELECT teacher_id, absence_date FROM teacher_absences "
            "WHERE absence_date BETWEEN %s AND %s", (ws, we))
        absent = {(a["teacher_id"], _as_date(a["absence_date"])) for a in cur.fetchall()}
        avail = [{t for t in teachers if (t, d) not in absent} for d in days]

        rooms = _load_rooms(cur)
        max_cap = max(r["capacity"] for r in rooms)

        # another round must not share this week
        cur.execute(
            "SELECT COUNT(*) AS c FROM exam_schedule "
            "WHERE exam_round <> %s AND exam_date BETWEEN %s AND %s",
            (body.exam_round, ws, we))
        if cur.fetchone()["c"] > 0:
            raise HTTPException(
                400, "Another exam round is already scheduled in that week. Choose another week.")

        # ---- batches and papers (labs and practicals included) ----
        cur.execute("SELECT batch_id, batch_name, department_id, semester FROM batches ORDER BY batch_id")
        batches = cur.fetchall()
        if not batches:
            raise HTTPException(400, "No batches found.")
        cur.execute("SELECT batch_id, COUNT(*) AS c FROM students GROUP BY batch_id")
        counts = {r["batch_id"]: r["c"] for r in cur.fetchall()}

        papers, batch_by_id, skipped = [], {}, []
        for b in batches:
            cur.execute(
                """SELECT * FROM subjects
                   WHERE department_id = %s AND semester = %s
                     AND (batch_id IS NULL OR batch_id = %s)
                     AND COALESCE(sessions_per_week, 0) > 0
                   ORDER BY subject_id""",
                (b["department_id"], b["semester"], b["batch_id"]))
            subs = cur.fetchall()
            if not subs:
                skipped.append(f"{b['batch_name']} (no subjects)")
                continue
            if len(subs) > 2 * D:
                raise HTTPException(
                    400, f"{b['batch_name']} has {len(subs)} subjects, but only {2 * D} papers "
                         f"fit in {D} exam day(s). Days skipped as holidays: "
                         f"{', '.join(skipped_days) or 'none'}.")
            size = max(counts.get(b["batch_id"], 0), 1)
            if size > max_cap:
                raise HTTPException(
                    400, f"{b['batch_name']} has {size} students but the largest room holds only {max_cap}.")
            batch_by_id[b["batch_id"]] = b
            for s in subs:
                if s.get("teacher_id"):
                    own.setdefault(s["subject_id"], set()).add(s["teacher_id"])
                papers.append({
                    "batch_id": b["batch_id"], "subject": s, "size": size,
                    "needs_lab": "lab" in str(s.get("room_type_required") or "").lower(),
                })
        if not papers:
            raise HTTPException(400, "No batch has subjects, so no exam can be scheduled.")
        P = len(papers)
        KS = range(len(SESSIONS))

        # ============ STAGE 1: papers -> (day, session), balanced ============
        m1 = cp_model.CpModel()
        x = {(i, d, k): m1.NewBoolVar(f"x_{i}_{d}_{k}")
             for i in range(P) for d in range(D) for k in KS}
        for i in range(P):
            m1.Add(sum(x[i, d, k] for d in range(D) for k in KS) == 1)

        by_batch = {}
        for i, p in enumerate(papers):
            by_batch.setdefault(p["batch_id"], []).append(i)

        # every batch gets its papers spread evenly over the days
        shortfalls = []
        for idxs in by_batch.values():
            n = len(idxs)
            floor_n = n // D
            for d in range(D):
                for k in KS:
                    m1.Add(sum(x[i, d, k] for i in idxs) <= 1)    # one paper at a time
                day_total = sum(x[i, d, k] for i in idxs for k in KS)
                m1.Add(day_total <= (1 if n <= D else 2))
                if floor_n > 0:                                   # "at least" is soft
                    sh = m1.NewIntVar(0, floor_n, f"sh_{len(shortfalls)}")
                    m1.Add(day_total + sh >= floor_n)
                    shortfalls.append(sh)

        # the whole department gets about the same number of papers every day
        day_load = [sum(x[i, d, k] for i in range(P) for k in KS) for d in range(D)]
        max_day = m1.NewIntVar(0, P, "max_day")
        min_day = m1.NewIntVar(0, P, "min_day")
        for d in range(D):
            m1.Add(day_load[d] <= max_day)
            m1.Add(day_load[d] >= min_day)

        sizes = sorted({p["size"] for p in papers})
        peak = m1.NewIntVar(0, P, "peak")
        for d in range(D):
            for k in KS:
                load = sum(x[i, d, k] for i in range(P))
                m1.Add(load <= len(avail[d]))                     # one teacher per room
                m1.Add(load <= peak)
                for sz in sizes:                                  # room capacity check
                    ge = [i for i in range(P) if papers[i]["size"] >= sz]
                    room_ge = sum(1 for r in rooms if r["capacity"] >= sz)
                    m1.Add(sum(x[i, d, k] for i in ge) <= room_ge)

        m1.Minimize(
            1000 * (max_day - min_day)
            + 200 * sum(shortfalls)
            + 20 * peak
            + sum(x[i, d, 1] for i in range(P) for d in range(D)))

        s1 = cp_model.CpSolver()
        s1.parameters.max_time_in_seconds = SOLVER_SECONDS
        s1.parameters.num_workers = 4
        st1 = s1.Solve(m1)
        if st1 not in (cp_model.OPTIMAL, cp_model.FEASIBLE):
            raise HTTPException(
                400,
                f"The {label} timetable could not be generated within one week (Mon-Fri) with the "
                f"current resources. {P} papers, {len(rooms)} rooms, {len(SESSIONS)} sessions per day, "
                f"teachers available per day: {[len(a) for a in avail]}. "
                "Suggestions: add rooms or teachers, make absent teachers available, "
                "or reduce the number of subjects.")

        slot_of = {}
        for i in range(P):
            for d in range(D):
                for k in KS:
                    if s1.Value(x[i, d, k]):
                        slot_of[i] = (d, k)
        slots = {}
        for i, sl in slot_of.items():
            slots.setdefault(sl, []).append(i)

        # ============ STAGE 2: invigilators, fair duty ============
        items = [{
            "date": days[slot_of[i][0]],
            "start": _mins(SESSIONS[slot_of[i][1]]["start"]),
            "end": _mins(SESSIONS[slot_of[i][1]]["end"]),
            "subject_id": papers[i]["subject"]["subject_id"],
        } for i in range(P)]
        inv_of, eligible = _assign_invigilators(items, teachers, absent, own)

        # ============ STAGE 3: rooms (labs first for lab papers) ============
        room_of = {}
        for idxs in slots.values():
            free = list(rooms)
            for i in sorted(idxs, key=lambda j: -papers[j]["size"]):
                fit = [r for r in free if r["capacity"] >= papers[i]["size"]]
                fit.sort(key=lambda r: (
                    0 if r["is_lab"] == papers[i]["needs_lab"] else 1,
                    r["capacity"], r["room_id"]))
                if not fit:
                    raise HTTPException(
                        400, "Could not find a free room that is big enough for one of the papers. "
                             "Add a bigger room or reduce the batch size.")
                room = fit[0]
                free.remove(room)
                room_of[i] = room

        # ============ SAVE (replace this round only) ============
        cur.execute("DELETE FROM exam_schedule WHERE exam_round = %s", (body.exam_round,))
        w_cur = conn.cursor()
        for i, p in enumerate(papers):
            d, k = slot_of[i]
            b = batch_by_id[p["batch_id"]]
            w_cur.execute(
                """INSERT INTO exam_schedule
                   (exam_name, exam_type, exam_round, batch_id, subject_id, semester, exam_date,
                    session_no, start_time, end_time, room_id, invigilator_id)
                   VALUES (%s,'internal',%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
                (f"{label} - {p['subject']['subject_name']}", body.exam_round,
                 p["batch_id"], p["subject"]["subject_id"], b["semester"], days[d],
                 SESSIONS[k]["no"], SESSIONS[k]["start"], SESSIONS[k]["end"],
                 room_of[i]["room_id"], inv_of[i]))
        conn.commit()
        w_cur.close()

        duty = {t: 0 for t in teachers}
        for t in inv_of:
            duty[t] += 1
        double_days = []
        for bid, idxs in by_batch.items():
            per_batch_day = {}
            for i in idxs:
                per_batch_day[slot_of[i][0]] = per_batch_day.get(slot_of[i][0], 0) + 1
            for d, c in per_batch_day.items():
                if c == 2:
                    double_days.append(
                        {"batch": batch_by_id[bid]["batch_name"], "date": days[d].isoformat()})

        per_day = {d.isoformat(): 0 for d in days}
        for i in range(P):
            per_day[days[slot_of[i][0]].isoformat()] += 1
        conflict_report = _find_conflicts(cur, body.exam_round)

        return {
            "success": True,
            "message": f"{label} timetable generated for {len(by_batch)} batch(es).",
            "exam_round": body.exam_round,
            "round_label": label,
            "total_exams": P,
            "start_date": days[0].isoformat(),
            "end_date": days[-1].isoformat(),
            "exam_days": [d.isoformat() for d in days],
            "double_paper_days": double_days,
            "fairness": _fairness(duty, teachers, eligible),
            "skipped_batches": skipped,
            "per_day": [{"date": k, "exams": v} for k, v in sorted(per_day.items())],
            "conflicts": conflict_report["conflicts"],
        }

    except HTTPException:
        conn.rollback()
        raise
    except Exception as e:
        conn.rollback()
        raise HTTPException(400, str(e))
    finally:
        cur.close()
        conn.close()


@router.post("/generate")
async def generate_internal_timetable(body: GenerateIn,
                                      user: dict = Depends(require_role("admin"))):
    result = await run_in_threadpool(_generate_sync, body)
    await _notify(
        f"{result['round_label']} Timetable Published",
        f"{result['round_label']}: {result['start_date']} to {result['end_date']}. "
        "Regular classes are suspended in this week. Check your exam schedule.")
    return result


# ------------------------------------------------------------
# EXTERNAL: import dates from a file, then assign room + invigilator
# ------------------------------------------------------------
def _parse_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d-%b-%Y", "%d %b %Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _parse_time(v):
    if isinstance(v, datetime):
        return v.time().replace(second=0, microsecond=0)
    if isinstance(v, dtime):
        return v.replace(second=0, microsecond=0)
    s = str(v or "").strip().upper()
    for fmt in ("%H:%M", "%H:%M:%S", "%I:%M %p", "%I:%M%p", "%I %p"):
        try:
            return datetime.strptime(s, fmt).time()
        except ValueError:
            pass
    return None


def _blank(v):
    return v is None or str(v).strip() == ""


def _read_rows(filename, data):
    name = (filename or "").lower()
    if name.endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        raw = [list(r) for r in csv.reader(io.StringIO(text))]
    elif name.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise HTTPException(500, "openpyxl is not installed. Run: pip install openpyxl")
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            raw = [list(r) for r in wb.active.iter_rows(values_only=True)]
            wb.close()
        except Exception as e:
            raise HTTPException(400, f"Could not read the Excel file: {e}")
    else:
        raise HTTPException(400, "Only CSV and XLSX files are supported.")

    if not raw:
        raise HTTPException(400, "The file is empty.")
    headers = [str(h or "").strip().lower() for h in raw[0]]
    missing = [c for c in EXTERNAL_REQUIRED if c not in headers]
    if missing:
        raise HTTPException(400, f"Missing required column(s): {', '.join(missing)}")
    idx = {h: i for i, h in enumerate(headers)}

    rows = []
    for n, r in enumerate(raw[1:], start=2):
        if not any(not _blank(c) for c in r):
            continue
        row = {"row": n}
        for c in EXTERNAL_COLUMNS:
            i = idx.get(c)
            row[c] = r[i] if i is not None and i < len(r) else None
        rows.append(row)
    return rows


def _import_external_sync(filename, data):
    rows = _read_rows(filename, data)
    if not rows:
        raise HTTPException(400, "No exam rows found in the file.")

    conn = get_db_connection()
    cur = conn.cursor(dictionary=True)
    try:
        cur.execute("SELECT batch_id, batch_name, department_id, semester FROM batches")
        batches = {str(b["batch_name"]).strip().lower(): b for b in cur.fetchall()}
        cur.execute("SELECT batch_id, COUNT(*) AS c FROM students GROUP BY batch_id")
        counts = {r["batch_id"]: r["c"] for r in cur.fetchall()}
        teachers, own = _load_teachers_and_own(cur)
        rooms = _load_rooms(cur)
        room_by_name = {r["room_name"].strip().lower(): r for r in rooms}

        cur.execute("SELECT DISTINCT exam_date FROM exam_schedule WHERE exam_round <> 'external'")
        internal_dates = {_as_date(r["exam_date"]) for r in cur.fetchall()}

        subj_cache = {}

        def subjects_of(b):
            key = b["batch_id"]
            if key not in subj_cache:
                cur.execute(
                    "SELECT * FROM subjects WHERE department_id = %s AND semester = %s "
                    "AND (batch_id IS NULL OR batch_id = %s)",
                    (b["department_id"], b["semester"], b["batch_id"]))
                subj_cache[key] = cur.fetchall()
            return subj_cache[key]

        errors, exams, seen = [], [], set()
        for r in rows:
            errs = []
            bname = str(r["batch_name"] or "").strip()
            b = batches.get(bname.lower())
            if not b:
                errs.append(f"batch '{bname}' does not exist")

            code = str(r["subject_code"] or "").strip()
            subject = None
            if b:
                for s in subjects_of(b):
                    if str(s.get("subject_code") or "").strip().lower() == code.lower() \
                            or str(s.get("subject_name") or "").strip().lower() == code.lower():
                        subject = s
                        break
                if not subject:
                    errs.append(f"subject '{code}' not found in {b['batch_name']}")

            d = _parse_date(r["exam_date"])
            if not d:
                errs.append("invalid exam_date (use YYYY-MM-DD or DD-MM-YYYY)")
            elif d.weekday() == 6:
                errs.append("exam_date is a Sunday")
            elif d in internal_dates:
                errs.append("internal exams are already scheduled on this date")

            st = DEFAULT_EXT_START if _blank(r["start_time"]) else _parse_time(r["start_time"])
            en = DEFAULT_EXT_END if _blank(r["end_time"]) else _parse_time(r["end_time"])
            if st is None:
                errs.append("invalid start_time")
            if en is None:
                errs.append("invalid end_time")
            if st and en and en <= st:
                errs.append("end_time must be after start_time")

            room = None
            if not _blank(r["room_name"]):
                room = room_by_name.get(str(r["room_name"]).strip().lower())
                if not room:
                    errs.append(f"room '{r['room_name']}' does not exist")

            if b and subject:
                key = (b["batch_id"], subject["subject_id"])
                if key in seen:
                    errs.append("this batch and subject appear twice in the file")
                seen.add(key)

            size = max(counts.get(b["batch_id"], 0), 1) if b else 1
            if room and room["capacity"] < size:
                errs.append(f"room '{room['room_name']}' holds {room['capacity']} but the batch has {size} students")

            if errs:
                errors.append(f"Row {r['row']}: " + "; ".join(errs))
                continue
            exams.append({
                "row": r["row"], "batch": b, "subject": subject, "date": d,
                "start": st, "end": en,
                "s": st.hour * 60 + st.minute, "e": en.hour * 60 + en.minute,
                "size": size, "room": room,
                "needs_lab": "lab" in str(subject.get("room_type_required") or "").lower(),
            })

        # a batch cannot have two overlapping exams
        for i in range(len(exams)):
            for j in range(i + 1, len(exams)):
                a, c = exams[i], exams[j]
                if a["batch"]["batch_id"] == c["batch"]["batch_id"] and a["date"] == c["date"] \
                        and _overlap(a["s"], a["e"], c["s"], c["e"]):
                    errors.append(
                        f"Row {c['row']}: {c['batch']['batch_name']} already has an exam at that time (row {a['row']})")

        # ---- rooms ----
        busy = {}

        def room_free(rid, d, s, e):
            return all(not (bd == d and _overlap(s, e, bs, be)) for bd, bs, be in busy.get(rid, []))

        if not errors:
            for ex in exams:
                if ex["room"]:
                    if not room_free(ex["room"]["room_id"], ex["date"], ex["s"], ex["e"]):
                        errors.append(
                            f"Row {ex['row']}: room '{ex['room']['room_name']}' is already used at that time")
                    else:
                        busy.setdefault(ex["room"]["room_id"], []).append((ex["date"], ex["s"], ex["e"]))
            for ex in sorted((e for e in exams if not e["room"]),
                             key=lambda e: (e["date"], e["s"], -e["size"])):
                fit = [r for r in rooms
                       if r["capacity"] >= ex["size"] and room_free(r["room_id"], ex["date"], ex["s"], ex["e"])]
                fit.sort(key=lambda r: (0 if r["is_lab"] == ex["needs_lab"] else 1, r["capacity"], r["room_id"]))
                if not fit:
                    errors.append(
                        f"Row {ex['row']}: no free room for {ex['batch']['batch_name']} on {ex['date']}")
                    continue
                ex["room"] = fit[0]
                busy.setdefault(fit[0]["room_id"], []).append((ex["date"], ex["s"], ex["e"]))

        if errors:
            shown = errors[:20]
            more = f"\n- ...and {len(errors) - 20} more" if len(errors) > 20 else ""
            raise HTTPException(
                400, "The file has problems, so nothing was imported:\n- " + "\n- ".join(shown) + more)

        # ---- invigilators ----
        d_min = min(e["date"] for e in exams)
        d_max = max(e["date"] for e in exams)
        cur.execute(
            "SELECT teacher_id, absence_date FROM teacher_absences WHERE absence_date BETWEEN %s AND %s",
            (d_min, d_max))
        absent = {(a["teacher_id"], _as_date(a["absence_date"])) for a in cur.fetchall()}
        for ex in exams:
            if ex["subject"].get("teacher_id"):
                own.setdefault(ex["subject"]["subject_id"], set()).add(ex["subject"]["teacher_id"])
        items = [{"date": e["date"], "start": e["s"], "end": e["e"],
                  "subject_id": e["subject"]["subject_id"]} for e in exams]
        inv_of, eligible = _assign_invigilators(items, teachers, absent, own)

        # ---- save (replace the old external schedule) ----
        cur.execute("DELETE FROM exam_schedule WHERE exam_round = 'external'")
        w = conn.cursor()
        for ex, tid in zip(exams, inv_of):
            w.execute(
                """INSERT INTO exam_schedule
                   (exam_name, exam_type, exam_round, batch_id, subject_id, semester, exam_date,
                    session_no, start_time, end_time, room_id, invigilator_id)
                   VALUES (%s,'external','external',%s,%s,%s,%s,1,%s,%s,%s,%s)""",
                (f"External Exam - {ex['subject']['subject_name']}",
                 ex["batch"]["batch_id"], ex["subject"]["subject_id"], ex["batch"]["semester"],
                 ex["date"], ex["start"].strftime("%H:%M:%S"), ex["end"].strftime("%H:%M:%S"),
                 ex["room"]["room_id"], tid))
        conn.commit()
        w.close()

        duty = {t: 0 for t in teachers}
        for t in inv_of:
            duty[t] += 1

        per_day = {}
        for e in exams:
            key = e["date"].isoformat()
            per_day[key] = per_day.get(key, 0) + 1
        conflict_report = _find_conflicts(cur, "external")

        return {
            "success": True,
            "message": f"{len(exams)} external exam(s) imported and invigilators assigned.",
            "exam_round": "external",
            "round_label": "External Exam",
            "total_exams": len(exams),
            "start_date": d_min.isoformat(),
            "end_date": d_max.isoformat(),
            "fairness": _fairness(duty, teachers, eligible),
            "per_day": [{"date": k, "exams": v} for k, v in sorted(per_day.items())],
            "conflicts": conflict_report["conflicts"],
        }

    except HTTPException:
        conn.rollback()
        raise
    except Exception as e:
        conn.rollback()
        raise HTTPException(400, str(e))
    finally:
        cur.close()
        conn.close()


@router.get("/external/template")
def external_template(user: dict = Depends(require_role("admin"))):
    out = io.StringIO()
    wr = csv.writer(out)
    wr.writerow(EXTERNAL_COLUMNS)
    wr.writerow(["BCA 5 Sem", "BCA125", "2026-12-01", "10:00", "13:00", ""])
    wr.writerow(["BCA DS 5 Sem", "IDS019", "2026-12-01", "10:00", "13:00", ""])
    return StreamingResponse(
        io.BytesIO(out.getvalue().encode("utf-8")),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=external_exam_template.csv"})


@router.post("/external/import")
async def import_external(file: UploadFile = File(...),
                          user: dict = Depends(require_role("admin"))):
    data = await file.read()
    if not data:
        raise HTTPException(400, "The uploaded file is empty.")
    result = await run_in_threadpool(_import_external_sync, file.filename, data)
    await _notify(
        "External Exam Schedule Published",
        f"External exams: {result['start_date']} to {result['end_date']}. "
        "Regular classes are suspended in those weeks. Check your exam schedule.")
    return result


# ------------------------------------------------------------
# read
# ------------------------------------------------------------
@router.get("/week-status")
def week_status(user: dict = Depends(get_current_user)):
    conn = get_db_connection()
    cur = conn.cursor(dictionary=True)
    try:
        w = exam_week_containing(cur, date.today())
        if not w:
            return {"active": False, "week_start": None, "week_end": None}
        return {"active": True, "week_start": str(w["week_start"]), "week_end": str(w["week_end"])}
    finally:
        cur.close()
        conn.close()


@router.get("/schedule")
def get_schedule(exam_round: Optional[str] = None,
                 user: dict = Depends(get_current_user)):
    """Admin: everything. Student: own batch only. Teacher: only own invigilation duties."""
    conn = get_db_connection()
    cur = conn.cursor(dictionary=True)
    try:
        role = str(user.get("role", "")).lower()
        q = """
            SELECT e.exam_id, e.exam_name, e.exam_type, e.exam_round, e.exam_date, e.session_no,
                   e.start_time, e.end_time, e.semester, e.batch_id, b.batch_name,
                   s.subject_id, s.subject_name, s.subject_code,
                   r.room_name, r.room_type,
                   e.invigilator_id, t.teacher_name AS invigilator_name
            FROM exam_schedule e
            JOIN batches b ON e.batch_id = b.batch_id
            JOIN subjects s ON e.subject_id = s.subject_id
            JOIN rooms r ON e.room_id = r.room_id
            JOIN teachers t ON e.invigilator_id = t.teacher_id
            WHERE 1 = 1"""
        params = []
        if exam_round in ROUND_LABELS:
            q += " AND e.exam_round = %s"
            params.append(exam_round)

        if role == "student":
            cur.execute("SELECT batch_id FROM students WHERE student_id = %s",
                        (user.get("student_id"),))
            st = cur.fetchone()
            if not st:
                return []
            q += " AND e.batch_id = %s"
            params.append(st["batch_id"])
        elif role == "teacher":
            q += " AND e.invigilator_id = %s"
            params.append(user.get("teacher_id"))
        elif role != "admin":
            raise HTTPException(403, "Not allowed")

        q += " ORDER BY e.exam_date, e.start_time, b.batch_name"
        cur.execute(q, tuple(params))
        rows = cur.fetchall()
        for r in rows:
            r["exam_date"] = str(r["exam_date"])
            r["start_time"] = _hhmm(r["start_time"])
            r["end_time"] = _hhmm(r["end_time"])
            r["round_label"] = ROUND_LABELS.get(r["exam_round"], r["exam_round"])
        return rows
    finally:
        cur.close()
        conn.close()


@router.get("/duty-summary")
def duty_summary(exam_round: str, user: dict = Depends(require_role("admin"))):
    conn = get_db_connection()
    cur = conn.cursor(dictionary=True)
    try:
        cur.execute(
            """SELECT t.teacher_id, t.teacher_name,
                      (SELECT COUNT(*) FROM exam_schedule e
                        WHERE e.invigilator_id = t.teacher_id AND e.exam_round = %s) AS duties,
                      (SELECT COUNT(*) FROM exam_schedule e
                        WHERE e.invigilator_id = t.teacher_id) AS total_duties
               FROM teachers t
               ORDER BY duties DESC, t.teacher_name""", (exam_round,))
        return cur.fetchall()
    finally:
        cur.close()
        conn.close()


@router.get("/conflicts")
def check_conflicts(exam_round: Optional[str] = None,
                    user: dict = Depends(require_role("admin"))):
    conn = get_db_connection()
    cur = conn.cursor(dictionary=True)
    try:
        report = _find_conflicts(cur, exam_round if exam_round in ROUND_LABELS else None)
        return {"ok": len(report["conflicts"]) == 0, **report}
    finally:
        cur.close()
        conn.close()


# ------------------------------------------------------------
# admin edits
# ------------------------------------------------------------
@router.put("/{exam_id}/invigilator")
def change_invigilator(exam_id: int, body: InvigilatorIn,
                       user: dict = Depends(require_role("admin"))):
    conn = get_db_connection()
    cur = conn.cursor(dictionary=True)
    try:
        cur.execute("SELECT exam_date FROM exam_schedule WHERE exam_id = %s", (exam_id,))
        ex = cur.fetchone()
        if not ex:
            raise HTTPException(404, "Exam not found")

        cur.execute("SELECT teacher_name FROM teachers WHERE teacher_id = %s", (body.teacher_id,))
        t = cur.fetchone()
        if not t:
            raise HTTPException(404, "Teacher not found")

        cur.execute("SELECT 1 FROM teacher_absences WHERE teacher_id = %s AND absence_date = %s",
                    (body.teacher_id, ex["exam_date"]))
        if cur.fetchone():
            raise HTTPException(400, f"{t['teacher_name']} is absent on that day.")

        cur.execute(
            """SELECT 1 FROM exam_schedule o
               JOIN exam_schedule e ON e.exam_id = %s
               WHERE o.invigilator_id = %s AND o.exam_date = e.exam_date
                 AND o.exam_id <> e.exam_id
                 AND o.start_time < e.end_time AND o.end_time > e.start_time
               LIMIT 1""", (exam_id, body.teacher_id))
        if cur.fetchone():
            raise HTTPException(400, f"{t['teacher_name']} already has duty at that time.")

        w = conn.cursor()
        w.execute("UPDATE exam_schedule SET invigilator_id = %s WHERE exam_id = %s",
                  (body.teacher_id, exam_id))
        conn.commit()
        w.close()
        return {"message": "Invigilator changed"}
    except HTTPException:
        conn.rollback()
        raise
    finally:
        cur.close()
        conn.close()


@router.delete("/clear")
def clear_schedule(exam_round: str, user: dict = Depends(require_role("admin"))):
    if exam_round not in ROUND_LABELS:
        raise HTTPException(400, "Unknown exam round")
    conn = get_db_connection()
    cur = conn.cursor()
    try:
        cur.execute("DELETE FROM exam_schedule WHERE exam_round = %s", (exam_round,))
        removed = cur.rowcount
        conn.commit()
        return {"message": f"{removed} exam(s) removed"}
    finally:
        cur.close()
        conn.close()


# ------------------------------------------------------------
# EXPORT (Excel / PDF) - admin only
# ------------------------------------------------------------
def _fmt12(hhmm):
    h, m = [int(p) for p in hhmm.split(":")[:2]]
    return f"{h % 12 or 12}:{m:02d} {'AM' if h < 12 else 'PM'}"


def _pretty(iso):
    return date.fromisoformat(iso).strftime("%a, %d %b %Y")


def _short(iso):
    return date.fromisoformat(iso).strftime("%a %d %b")


def _time_range(r):
    return f"{_fmt12(r['start_time'])} - {_fmt12(r['end_time'])}"


def _room_text(r):
    return f"{r['room_name']} ({r['room_type']})" if r.get("room_type") else r["room_name"]


def _build_xlsx(rows, label):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.properties import PageSetupProperties
    except ImportError:
        raise HTTPException(500, "openpyxl is not installed. Run: pip install openpyxl")

    NAVY, LIGHT, ZEBRA = "1E3A5F", "EEF2F7", "F8FAFC"
    thin = Side(style="thin", color="CBD5E1")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    dates = sorted({r["exam_date"] for r in rows})
    batches = sorted({r["batch_name"] for r in rows})
    subtitle = (f"{_pretty(dates[0])} to {_pretty(dates[-1])}   |   "
                f"{len(rows)} papers   |   {len(batches)} batches")

    def sheet(ws, headers, data, widths):
        n = len(headers)
        banner = [(DEPARTMENT_NAME, 15, True, NAVY),
                  (f"{label} - Examination Timetable", 12, True, NAVY),
                  (subtitle, 10, False, LIGHT)]
        for ri, (text, size, bold, fill) in enumerate(banner, start=1):
            ws.merge_cells(start_row=ri, start_column=1, end_row=ri, end_column=n)
            c = ws.cell(ri, 1, text)
            c.font = Font(bold=bold, size=size, color="FFFFFF" if fill == NAVY else "334155")
            c.fill = PatternFill("solid", fgColor=fill)
            c.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[ri].height = 26 if ri < 3 else 20
        hr = 5
        for ci, h in enumerate(headers, start=1):
            c = ws.cell(hr, ci, h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="2C5282")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = box
        ws.row_dimensions[hr].height = 22
        for ri, row in enumerate(data, start=hr + 1):
            lines = 1
            for ci, v in enumerate(row, start=1):
                c = ws.cell(ri, ci, v)
                c.border = box
                c.alignment = Alignment(vertical="center", wrap_text=True)
                if (ri - hr) % 2 == 0:
                    c.fill = PatternFill("solid", fgColor=ZEBRA)
                lines = max(lines, str(v).count("\n") + 1)
            ws.row_dimensions[ri].height = max(20, 15 * lines)
        for ci, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(ci)].width = w
        ws.freeze_panes = ws.cell(hr + 1, 1)
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Timetable"
    sheet(ws1,
          ["Date", "Time", "Batch", "Code", "Subject", "Room / Lab", "Invigilator"],
          [[_pretty(r["exam_date"]), _time_range(r), r["batch_name"], r["subject_code"],
            r["subject_name"], _room_text(r), r["invigilator_name"]] for r in rows],
          [18, 20, 16, 12, 38, 22, 24])

    ws2 = wb.create_sheet("Batch-wise")
    grid = []
    for b in batches:
        line = [b]
        for d in dates:
            items = [r for r in rows if r["batch_name"] == b and r["exam_date"] == d]
            line.append("\n\n".join(
                f"{r['subject_name']}\n{_time_range(r)}\n{r['room_name']}" for r in items))
        grid.append(line)
    sheet(ws2, ["Batch"] + [_short(d) for d in dates], grid, [16] + [26] * len(dates))

    ws3 = wb.create_sheet("Invigilators")
    totals = {}
    for r in rows:
        totals[r["invigilator_name"]] = totals.get(r["invigilator_name"], 0) + 1
    duty = sorted(rows, key=lambda r: (r["invigilator_name"], r["exam_date"], r["start_time"]))
    sheet(ws3,
          ["Invigilator", "Total duties", "Date", "Time", "Batch", "Room / Lab"],
          [[r["invigilator_name"], totals[r["invigilator_name"]], _pretty(r["exam_date"]),
            _time_range(r), r["batch_name"], _room_text(r)] for r in duty],
          [26, 14, 18, 20, 16, 22])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build_pdf(rows, label):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.platypus import (PageBreak, Paragraph, SimpleDocTemplate,
                                        Spacer, Table, TableStyle)
    except ImportError:
        raise HTTPException(500, "reportlab is not installed. Run: pip install reportlab")

    NAVY = colors.HexColor("#1e3a5f")
    dates = sorted({r["exam_date"] for r in rows})
    batches = sorted({r["batch_name"] for r in rows})

    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Title"], fontSize=17, textColor=NAVY, spaceAfter=2)
    h2 = ParagraphStyle("h2", parent=ss["Normal"], fontSize=12, alignment=1,
                        textColor=NAVY, fontName="Helvetica-Bold", spaceAfter=2)
    sub = ParagraphStyle("sub", parent=ss["Normal"], fontSize=9.5, alignment=1,
                         textColor=colors.HexColor("#475569"), spaceAfter=10)
    sec = ParagraphStyle("sec", parent=ss["Normal"], fontSize=12, textColor=NAVY,
                         fontName="Helvetica-Bold", spaceBefore=4, spaceAfter=6)
    cell = ParagraphStyle("cell", parent=ss["Normal"], fontSize=8.5, leading=11)
    head = ParagraphStyle("head", parent=cell, textColor=colors.white, fontName="Helvetica-Bold")

    def P(text, style=cell):
        return Paragraph(escape(str(text)).replace("\n", "<br/>"), style)

    def table(data, widths):
        t = Table([[P(c, head) for c in data[0]]] + [[P(c) for c in row] for row in data[1:]],
                  colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), NAVY),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        return t

    def header():
        return [Paragraph(escape(DEPARTMENT_NAME), h1),
                Paragraph(escape(f"{label} - Examination Timetable"), h2),
                Paragraph(escape(f"{_pretty(dates[0])} to {_pretty(dates[-1])}  |  "
                                 f"{len(rows)} papers  |  {len(batches)} batches"), sub)]

    story = header()
    story.append(Paragraph("Date-wise timetable", sec))
    story.append(table(
        [["Date", "Time", "Batch", "Subject", "Room / Lab", "Invigilator"]] +
        [[_pretty(r["exam_date"]), _time_range(r), r["batch_name"],
          f"{r['subject_name']} ({r['subject_code']})", _room_text(r), r["invigilator_name"]]
         for r in rows],
        [95, 105, 85, 215, 120, 174]))

    story.append(PageBreak())
    story += header()
    story.append(Paragraph("Batch-wise datesheet", sec))
    first_w = 90
    col_w = (794 - first_w) / max(len(dates), 1)
    grid = [["Batch"] + [_short(d) for d in dates]]
    for b in batches:
        line = [b]
        for d in dates:
            items = [r for r in rows if r["batch_name"] == b and r["exam_date"] == d]
            line.append("\n".join(f"{r['subject_name']}\n{_time_range(r)}" for r in items) or "-")
        grid.append(line)
    story.append(table(grid, [first_w] + [col_w] * len(dates)))

    story.append(PageBreak())
    story += header()
    story.append(Paragraph("Invigilation duty chart", sec))
    by_t = {}
    for r in rows:
        by_t.setdefault(r["invigilator_name"], []).append(r)
    duty = [["Invigilator", "Duties", "Duty details"]]
    for name in sorted(by_t):
        items = sorted(by_t[name], key=lambda r: (r["exam_date"], r["start_time"]))
        duty.append([name, len(items), "\n".join(
            f"{_short(r['exam_date'])}, {_fmt12(r['start_time'])} - {r['room_name']} - {r['batch_name']}"
            for r in items)])
    story.append(table(duty, [150, 55, 589]))

    story.append(Spacer(1, 40))
    sig = Table([["Controller of Examinations", "Head of Department"]], colWidths=[397, 397])
    sig.setStyle(TableStyle([
        ("ALIGN", (0, 0), (-1, -1), "CENTER"), ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("LINEABOVE", (0, 0), (0, 0), 0.6, colors.grey), ("LINEABOVE", (1, 0), (1, 0), 0.6, colors.grey),
    ]))
    story.append(sig)

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24,
                      topMargin=24, bottomMargin=24,
                      title=f"{label} Timetable").build(story)
    return buf.getvalue()


@router.get("/export")
def export_schedule(exam_round: str,
                    file_format: str = Query("xlsx", alias="format"),
                    user: dict = Depends(require_role("admin"))):
    if exam_round not in ROUND_LABELS:
        raise HTTPException(400, "Unknown exam round")
    if file_format not in ("xlsx", "pdf"):
        raise HTTPException(400, "Format must be xlsx or pdf")
    rows = get_schedule(exam_round=exam_round, user=user)
    if not rows:
        raise HTTPException(404, "There is no schedule to download for this round.")

    label = ROUND_LABELS[exam_round]
    if file_format == "xlsx":
        data = _build_xlsx(rows, label)
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        data = _build_pdf(rows, label)
        media = "application/pdf"
    return StreamingResponse(
        io.BytesIO(data), media_type=media,
        headers={"Content-Disposition":
                 f"attachment; filename=exam_timetable_{exam_round}.{file_format}"})