"""
SchedAI - import_logic.py

Pure logic for bulk import (no database, no FastAPI), so it is easy to test:
  - read an Excel (.xlsx) or CSV file into rows
  - validate student rows and syllabus (subject) rows
  - build a ready-to-fill Excel template

The router in imports.py loads the database data into a `ctx` dictionary,
calls these functions, and saves the valid rows.
"""

import csv
import io
import re
from datetime import date, datetime, time as dtime

from openpyxl import Workbook, load_workbook

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


class ImportFileError(Exception):
    """The uploaded file itself is unusable (wrong type, missing columns...)."""


# ------------------------------------------------------------------
# Columns
# ------------------------------------------------------------------
STUDENT_COLUMNS = [
    "department", "batch_name", "semester", "section",
    "roll_number", "student_name", "email",
]
STUDENT_REQUIRED = ["batch_name", "semester", "roll_number", "student_name", "email"]

SYLLABUS_COLUMNS = [
    "department", "batch_name", "semester", "subject_code",
    "subject_name", "sessions_per_week", "type", "teacher_email",
]
SYLLABUS_REQUIRED = ["semester", "subject_name"]

CALENDAR_COLUMNS = ["type", "name", "date", "end_date", "time", "location", "description"]
CALENDAR_REQUIRED = ["type", "name", "date"]

REQUIRED = {"students": STUDENT_REQUIRED, "syllabus": SYLLABUS_REQUIRED, "calendar": CALENDAR_REQUIRED}
COLUMNS = {"students": STUDENT_COLUMNS, "syllabus": SYLLABUS_COLUMNS, "calendar": CALENDAR_COLUMNS}

# aliases used only for the calendar file ("name" means the holiday/event name here)
CALENDAR_ALIASES = {
    "kind": "type", "category": "type", "holiday_event": "type",
    "holiday_name": "name", "event_name": "name", "title": "name", "holiday": "name",
    "event": "name", "holiday_event_name": "name", "occasion": "name", "festival": "name",
    "name_of_holiday": "name", "name_of_event": "name",
    "holiday_date": "date", "event_date": "date", "start_date": "date", "from_date": "date",
    "from": "date", "on": "date",
    "to_date": "end_date", "end": "end_date", "till": "end_date", "to": "end_date",
    "event_time": "time", "start_time": "time",
    "venue": "location", "place": "location",
    "details": "description", "remarks": "description", "note": "description", "notes": "description",
}

# different names people use in their sheets -> our column name
ALIASES = {
    "dept": "department", "department_name": "department", "course": "department",
    "batch": "batch_name", "batchname": "batch_name", "class": "batch_name",
    "sem": "semester", "semester_no": "semester",
    "roll": "roll_number", "roll_no": "roll_number", "rollno": "roll_number",
    "rollnumber": "roll_number", "enrollment_no": "roll_number",
    "name": "student_name", "student": "student_name", "studentname": "student_name",
    "full_name": "student_name",
    "email_id": "email", "e_mail": "email", "mail": "email", "email_address": "email",
    "code": "subject_code", "subjectcode": "subject_code", "paper_code": "subject_code",
    "subject": "subject_name", "subjectname": "subject_name", "paper": "subject_name",
    "paper_name": "subject_name", "subject_title": "subject_name",
    "sessions": "sessions_per_week", "classes_per_week": "sessions_per_week",
    "lectures_per_week": "sessions_per_week", "periods_per_week": "sessions_per_week",
    "subject_type": "type", "lecture_lab": "type", "theory_lab": "type",
    "teacher": "teacher_email", "teacher_emails": "teacher_email",
    "faculty_email": "teacher_email",
}


def _norm_header(value, kind=None):
    s = str(value or "").strip().lower()
    s = re.sub(r"[\s/\-\.]+", "_", s)
    s = re.sub(r"_+", "_", s).strip("_")
    if kind == "calendar":
        return CALENDAR_ALIASES.get(s, s)
    return ALIASES.get(s, s)


def _cell(value):
    """Cell -> clean string (5.0 -> '5', None -> '')."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.hour == 0 and value.minute == 0 and value.second == 0:
            return value.date().isoformat()
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, dtime):
        return value.strftime("%H:%M")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


# ------------------------------------------------------------------
# Reading the file
# ------------------------------------------------------------------
def read_table(filename, data, kind):
    """
    Returns a list of dicts: {"_row": <sheet row number>, <column>: <text>, ...}
    Raises ImportFileError for an unusable file.
    """
    name = (filename or "").lower()
    required = REQUIRED[kind]

    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:
            raise ImportFileError(f"Could not read the Excel file: {e}")
        ws = wb.active
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
    elif name.endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        grid = [list(r) for r in csv.reader(io.StringIO(text))]
    else:
        raise ImportFileError("Only Excel (.xlsx) or CSV (.csv) files are supported.")

    # first non-empty row = header
    header_idx = next(
        (i for i, r in enumerate(grid) if any(_cell(c) for c in r)), None)
    if header_idx is None:
        raise ImportFileError("The file is empty.")

    headers = [_norm_header(c, kind) for c in grid[header_idx]]
    missing = [c for c in required if c not in headers]
    if missing:
        raise ImportFileError(
            "Missing required column(s): " + ", ".join(missing)
            + ". Download the template to see the expected columns.")

    rows = []
    for offset, raw in enumerate(grid[header_idx + 1:], start=header_idx + 2):
        if not any(_cell(c) for c in raw):
            continue                                   # blank row
        item = {"_row": offset}
        for col in COLUMNS[kind]:
            item[col] = _cell(raw[headers.index(col)]) if col in headers and headers.index(col) < len(raw) else ""
        rows.append(item)

    if not rows:
        raise ImportFileError("No data rows found below the header row.")
    if len(rows) > 5000:
        raise ImportFileError("Too many rows (limit is 5000 per file).")
    return rows


# ------------------------------------------------------------------
# Shared helpers
# ------------------------------------------------------------------
def _resolve_department(value, ctx, errors):
    departments = ctx["departments"]                    # {lower name: {"id","name"}}
    value = (value or "").strip()
    if not value:
        if len(departments) == 1:
            return next(iter(departments.values()))
        errors.append("Department is required")
        return None
    dept = departments.get(value.lower())
    if not dept:
        errors.append(f"Department '{value}' does not exist")
    return dept


def _parse_semester(value, errors):
    try:
        sem = int(float(value))
    except (TypeError, ValueError):
        errors.append("Semester must be a number")
        return None
    if sem < 1 or sem > 8:
        errors.append("Semester must be between 1 and 8")
        return None
    return sem


# ------------------------------------------------------------------
# STUDENTS
# ------------------------------------------------------------------
def validate_student_rows(rows, ctx):
    """
    ctx = {
      "departments":   {lower name: {"id", "name"}},
      "batches":       {(department_id, lower batch name): {"id", "semester"}},
      "existing_rolls":  set of lower roll numbers already in the database,
      "existing_emails": set of lower emails already in the database,
    }
    """
    results = []
    seen_rolls, seen_emails = set(), set()
    new_batch_semester = {}      # (dept_id, lower name) -> semester, for new batches in this file

    for row in rows:
        errors = []

        dept = _resolve_department(row["department"], ctx, errors)
        batch_name = row["batch_name"].strip()
        roll = row["roll_number"].strip()
        student = row["student_name"].strip()
        email = row["email"].strip().lower()
        section = row["section"].strip() or None
        semester = _parse_semester(row["semester"], errors)

        if not batch_name:
            errors.append("Batch name is required")
        if not roll:
            errors.append("Roll number is required")
        if not student:
            errors.append("Student name is required")
        if not email:
            errors.append("Email is required")
        elif not EMAIL_RE.match(email):
            errors.append("Invalid email format")

        # duplicates
        if roll:
            if roll.lower() in seen_rolls:
                errors.append("Duplicate roll number in this file")
            elif roll.lower() in ctx["existing_rolls"]:
                errors.append(f"Roll number '{roll}' already exists")
            seen_rolls.add(roll.lower())
        if email:
            if email in seen_emails:
                errors.append("Duplicate email in this file")
            elif email in ctx["existing_emails"]:
                errors.append(f"Email '{email}' already exists")
            seen_emails.add(email)

        # batch
        batch_id, is_new_batch = None, False
        if dept and batch_name and semester is not None:
            key = (dept["id"], batch_name.lower())
            existing = ctx["batches"].get(key)
            if existing:
                batch_id = existing["id"]
                if int(existing["semester"]) != semester:
                    errors.append(
                        f"Batch '{batch_name}' is semester {existing['semester']}, "
                        f"but this row says semester {semester}")
            else:
                is_new_batch = True
                prev = new_batch_semester.get(key)
                if prev is not None and prev != semester:
                    errors.append(
                        f"Batch '{batch_name}' appears with two different semesters in this file")
                new_batch_semester.setdefault(key, semester)

        results.append({
            "row": row["_row"],
            "department_id": dept["id"] if dept else None,
            "department": dept["name"] if dept else row["department"],
            "batch_name": batch_name,
            "batch_id": batch_id,
            "new_batch": is_new_batch,
            "semester": semester,
            "section": section,
            "roll_number": roll,
            "student_name": student,
            "email": email,
            "valid": not errors,
            "errors": errors,
        })

    return results


# ------------------------------------------------------------------
# SYLLABUS (subjects)
# ------------------------------------------------------------------
def validate_syllabus_rows(rows, ctx):
    """
    ctx = {
      "departments":       {lower name: {"id", "name"}},
      "batches":           {(department_id, lower batch name): {"id", "semester"}},
      "subjects_by_code":  {lower subject_code: subject_id},
      "subjects_by_key":   {(department_id, lower name, semester, batch_id): subject_id},
      "teachers_by_email": {lower email: teacher_id},
    }
    """
    results = []
    seen_codes, seen_keys = set(), set()

    for row in rows:
        errors = []

        dept = _resolve_department(row["department"], ctx, errors)
        semester = _parse_semester(row["semester"], errors)
        code = row["subject_code"].strip()
        name = row["subject_name"].strip()
        batch_name = row["batch_name"].strip()

        if not name:
            errors.append("Subject name is required")

        # sessions per week
        sessions = 3
        raw_sessions = row["sessions_per_week"].strip()
        if raw_sessions:
            try:
                sessions = int(float(raw_sessions))
                if sessions < 0 or sessions > 10:
                    errors.append("Sessions per week must be between 0 and 10")
            except ValueError:
                errors.append("Sessions per week must be a number")

        # theory / lab
        room_type_required = None
        kind = row["type"].strip().lower()
        if kind in ("lab", "practical", "laboratory"):
            room_type_required = "lab"
        elif kind in ("", "theory", "lecture", "classroom"):
            room_type_required = None
        else:
            errors.append(f"Type '{row['type']}' is not valid (use Theory or Lab)")

        # batch (optional: blank = every batch of this department + semester)
        batch_id = None
        if batch_name:
            if dept:
                found = ctx["batches"].get((dept["id"], batch_name.lower()))
                if not found:
                    errors.append(f"Batch '{batch_name}' does not exist (import its students first)")
                else:
                    batch_id = found["id"]
                    if semester is not None and int(found["semester"]) != semester:
                        errors.append(
                            f"Batch '{batch_name}' is semester {found['semester']}, "
                            f"but this row says semester {semester}")

        # teachers (optional, comma separated emails)
        teacher_ids = []
        raw_teachers = row["teacher_email"].strip()
        if raw_teachers:
            for mail in re.split(r"[,;\n]+", raw_teachers):
                mail = mail.strip().lower()
                if not mail:
                    continue
                tid = ctx["teachers_by_email"].get(mail)
                if tid is None:
                    errors.append(f"No teacher with email '{mail}'")
                elif tid not in teacher_ids:
                    teacher_ids.append(tid)

        # create or update?
        action, existing_id = "create", None
        if dept:
            if code and code.lower() in ctx["subjects_by_code"]:
                action, existing_id = "update", ctx["subjects_by_code"][code.lower()]
            elif not code and semester is not None:
                key = (dept["id"], name.lower(), semester, batch_id)
                if key in ctx["subjects_by_key"]:
                    action, existing_id = "update", ctx["subjects_by_key"][key]

        # duplicates inside the file
        if code:
            if code.lower() in seen_codes:
                errors.append(f"Subject code '{code}' appears twice in this file")
            seen_codes.add(code.lower())
        elif dept and semester is not None and name:
            key = (dept["id"], name.lower(), semester, batch_id)
            if key in seen_keys:
                errors.append("This subject appears twice in this file")
            seen_keys.add(key)

        results.append({
            "row": row["_row"],
            "department_id": dept["id"] if dept else None,
            "department": dept["name"] if dept else row["department"],
            "batch_name": batch_name,
            "batch_id": batch_id,
            "semester": semester,
            "subject_code": code or None,
            "subject_name": name,
            "sessions_per_week": sessions,
            "room_type_required": room_type_required,
            "type": "Lab" if room_type_required else "Theory",
            "teacher_ids": teacher_ids,
            "action": action,
            "existing_id": existing_id,
            "valid": not errors,
            "errors": errors,
        })

    return results


# ------------------------------------------------------------------
# CALENDAR (holidays + events)
# ------------------------------------------------------------------
_DATE_FORMATS = (
    "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y",
    "%d %b %Y", "%d %B %Y", "%b %d, %Y", "%B %d, %Y",
)
_TIME_FORMATS = ("%H:%M", "%H:%M:%S", "%I:%M %p", "%I:%M%p", "%I %p", "%H.%M")


def _parse_date(value, errors, label="Date"):
    v = (value or "").strip()
    if not v:
        return None
    if re.match(r"^\d{4}-\d{2}-\d{2}[ T]", v):
        v = v[:10]
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(v, fmt).date().isoformat()
        except ValueError:
            continue
    errors.append(f"{label} '{value}' is not a valid date (use YYYY-MM-DD or DD/MM/YYYY)")
    return None


def _parse_time(value, errors):
    v = (value or "").strip()
    if not v:
        return None
    if re.match(r"^\d{4}-\d{2}-\d{2}[ T]", v):
        v = v[11:]
    for fmt in _TIME_FORMATS:
        try:
            return datetime.strptime(v.upper(), fmt).strftime("%H:%M:00")
        except ValueError:
            continue
    errors.append(f"Time '{value}' is not valid (use e.g. 10:30 or 2:00 PM)")
    return None


def validate_calendar_rows(rows, ctx):
    """
    ctx = {
      "holidays": set of (lower name, 'YYYY-MM-DD') already in the database,
      "events":   set of (lower name, 'YYYY-MM-DD') already in the database,
    }
    """
    results = []
    seen = set()

    for row in rows:
        errors = []

        raw_type = row["type"].strip().lower()
        if raw_type in ("holiday", "h", "hol", "vacation", "leave"):
            kind = "holiday"
        elif raw_type in ("event", "e", "seminar", "workshop", "guest lecture",
                          "department event", "fest", "function"):
            kind = "event"
        else:
            kind = None
            errors.append(f"Type '{row['type']}' is not valid (use Holiday or Event)")

        name = row["name"].strip()
        if not name:
            errors.append("Name is required")

        day = _parse_date(row["date"], errors, "Date")
        if not row["date"].strip():
            errors.append("Date is required")

        end_day = None
        if row["end_date"].strip():
            end_day = _parse_date(row["end_date"], errors, "End date")
            if day and end_day and end_day < day:
                errors.append("End date is before the start date")
            if kind == "event":
                end_day = None                      # events are single-day here

        event_time = None
        if row["time"].strip() and kind != "holiday":
            event_time = _parse_time(row["time"], errors)

        # duplicates (inside the file, and already in the database)
        if kind and name and day:
            key = (kind, name.lower(), day)
            if key in seen:
                errors.append("This entry appears twice in this file")
            seen.add(key)

            existing = ctx["holidays"] if kind == "holiday" else ctx["events"]
            if (name.lower(), day) in existing:
                errors.append(f"This {kind} already exists in the calendar")

        results.append({
            "row": row["_row"],
            "type": kind,
            "name": name,
            "date": day,
            "end_date": end_day,
            "time": event_time,
            "location": row["location"].strip() or None,
            "description": row["description"].strip() or None,
            "valid": not errors,
            "errors": errors,
        })

    return results


# ------------------------------------------------------------------
# Excel templates
# ------------------------------------------------------------------
def make_template(kind):
    wb = Workbook()
    ws = wb.active

    if kind == "students":
        ws.title = "Students"
        ws.append(STUDENT_COLUMNS)
        ws.append(["Computer Applications", "BCA 5 Sem", 5, "A", "BCA5001", "Aarav Sharma", "aarav.sharma@example.com"])
        ws.append(["Computer Applications", "BCA 5 Sem", 5, "A", "BCA5002", "Diya Singh", "diya.singh@example.com"])
        ws.append(["Computer Applications", "BCA 3 Sem", 3, "A", "BCA3001", "Rohan Verma", "rohan.verma@example.com"])
        notes = wb.create_sheet("Notes")
        for line in [
            "One row per student. Several batches can be in the same file.",
            "department: exact department name (can be left blank if there is only one department).",
            "batch_name: if the batch does not exist yet it is created automatically.",
            "semester: 1 to 8, and it must match the batch.",
            "section: optional.",
            "email: the student will use THIS email to sign up, so it must be correct.",
        ]:
            notes.append([line])
    elif kind == "calendar":
        ws.title = "Calendar"
        ws.append(CALENDAR_COLUMNS)
        ws.append(["Holiday", "Diwali Break", "2026-11-08", "2026-11-12", "", "", "College closed for Diwali"])
        ws.append(["Holiday", "Gandhi Jayanti", "2026-10-02", "", "", "", ""])
        ws.append(["Event", "Guest Lecture on AI", "2026-10-15", "", "11:00 AM", "Seminar Hall", "By an industry expert"])
        notes = wb.create_sheet("Notes")
        for line in [
            "One row per holiday or event. Put both in the same file.",
            "type: Holiday or Event.",
            "date: YYYY-MM-DD or DD/MM/YYYY (a normal Excel date cell also works).",
            "end_date: only for holidays that last several days. Optional.",
            "time and location: for events only. Optional.",
            "Students and teachers are notified automatically after the import.",
        ]:
            notes.append([line])
    else:
        ws.title = "Syllabus"
        ws.append(SYLLABUS_COLUMNS)
        ws.append(["Computer Applications", "", 5, "BCA501", "Database Management System", 4, "Theory", "teacher1@example.com"])
        ws.append(["Computer Applications", "", 5, "BCA502", "DBMS Lab", 2, "Lab", "teacher1@example.com, teacher2@example.com"])
        ws.append(["Computer Applications", "BCA DS 5 Sem", 5, "BCADS501", "Data Visualization", 3, "Theory", ""])
        notes = wb.create_sheet("Notes")
        for line in [
            "One row per subject.",
            "batch_name: leave blank to apply the subject to EVERY batch of that department and semester.",
            "sessions_per_week: classes per week (0 = keep it out of the timetable). Default 3.",
            "type: Theory or Lab (Lab subjects are placed only in laboratories).",
            "teacher_email: who can teach it (separate several emails with commas). Optional.",
            "If a subject_code already exists, that subject is UPDATED instead of duplicated.",
        ]:
            notes.append([line])

    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 26

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()